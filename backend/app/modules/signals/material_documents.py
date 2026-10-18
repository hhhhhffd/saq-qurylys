"""In-memory DOCX, XLSX, and PDF formats for a signal-card materials package."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from re import compile
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from jinja2 import Environment, FileSystemLoader, select_autoescape
from weasyprint import HTML
from weasyprint.document import Document as PDFDocument

from app.xlsx_exports import ExportColumn, build_workbook

TEMPLATES_DIR = Path(__file__).resolve().parents[2] / "templates"
_INVALID_XML = compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_env = Environment(loader=FileSystemLoader(TEMPLATES_DIR), autoescape=select_autoescape(("html", "xml")))

_TEXT = {
    "ru": {
        "title": "Материалы к сигнальной карточке",
        "cover": "Сопроводительное письмо",
        "recipient": "Получатель",
        "card_number": "Номер сигнальной карточки",
        "formed_at": "Сформировано",
        "card_created_at": "Карточка создана",
        "snapshot_at": "Снимок данных",
        "object_code": "Код объекта",
        "object_name": "Наименование объекта",
        "developer": "Застройщик",
        "bin": "БИН",
        "address": "Адрес",
        "territory": "Территория",
        "zone": "Сигнальная зона (справочно)",
        "score": "Баллы (справочно)",
        "confirmed": "Подтверждено в снимке",
        "to_check": "Требует проверки",
        "no_items": "В этом разделе снимка нет перечисленных сведений.",
        "draft": "ЧЕРНОВИК · НЕ НАПРАВЛЕНО",
        "demo_stamp": "ДЕМО-ДАННЫЕ",
        "neutral_note": (
            "Сведения приведены по сохранённому снимку карточки и предназначены для проверки. "
            "Документ не устанавливает факт нарушения или виновность и сам по себе не подтверждает "
            "отправку карточки либо приложений. Сигнальная зона и баллы являются справочными индикаторами."
        ),
        "outgoing_number": "Исходящий номер",
        "signatory": "Подпись / Ф. И. О.",
        "date": "Дата",
        "appendices": "Перечень приложений",
        "appendix_no": "№",
        "appendix_name": "Наименование",
        "appendix_pages": "Листы",
        "no_data": "Нет данных в снимке",
        "publication_sheet": "Публикации",
        "publication_meta_sheet": "Сведения",
        "publication_title": "Публикации из снимка сигнальной карточки",
        "ad_id": "ID публикации",
        "platform": "Источник / платформа",
        "account": "Аккаунт",
        "published_at": "Дата публикации",
        "captured_at": "Дата захвата",
        "url": "URL из снимка",
        "ad_status": "Статус публикации в снимке",
        "fragments": "Фрагменты формулировок",
        "review_status": "Статус проверки формулировок в снимке",
        "appendix_refs": "Ссылки на приложения",
        "demo": "Демо-данные",
        "account_missing": "Не записан в снимке",
        "human_confirmed": "Подтверждено человеком",
        "human_rejected": "Отклонено человеком",
        "ai_hypothesis": "Гипотеза ИИ — не подтверждена человеком",
        "unreviewed": "Отдельная проверка человеком не отмечена в снимке",
        "fact_kind_note": "Классификация «факт» сама по себе не означает подтверждение человеком.",
        "snapshot_only": "Включены только публикации из сохранённого снимка карточки.",
        "metadata": "Параметр",
        "value": "Значение",
        "row_count": "Строк публикаций",
        "card_state": "Состояние карточки",
        "draft_state": "ЧЕРНОВИК · НЕ НАПРАВЛЕНО",
        "not_draft_state": "Не помечена как черновик в переданных данных",
        "package_note": "Этот файл сам по себе не подтверждает отправку карточки или пакета.",
        "snapshot_source": "Источник строк",
        "account_source": "Аккаунт-источник",
        "source_note": "Аккаунт-источник отсутствует в сохранённом снимке.",
        "review_rule": "Правило статуса проверки",
        "dispatch_status": "Сведения об отправке",
        "empty_rows": "В снимке нет публикаций.",
        "inventory_title": "Опись приложений",
        "inventory_intro": "Приложения к сигнальной карточке.",
        "page": "Стр.",
        "inventory_file": "Имя файла",
        "inventory_label": "Описание",
        "inventory_source": "Источник",
        "inventory_id": "ID доказательства",
        "inventory_type": "Тип файла",
        "inventory_size": "Размер, байт",
        "inventory_sha": "SHA-256",
        "inventory_capture": "Время захвата",
        "inventory_url": "Источник / URL",
        "inventory_included": "Включение в пакет",
        "included": "Включён",
        "omitted": "Не включён",
        "omission_title": "Исключения и недоступные файлы",
        "omission_reason": "Причина / примечание",
        "no_omissions": "Нет записанных исключений.",
        "missing_metadata": "Не зафиксировано",
        "self_exclusion": "Эта опись не содержит собственного размера или хеша.",
        "inventory_note": (
            "Опись перечисляет переданные метаданные и не устанавливает правовое значение файлов, "
            "факт нарушения или виновность."
        ),
    },
    "kk": {
        "title": "Сигналдық карточкаға арналған материалдар",
        "cover": "Ілеспе хат",
        "recipient": "Алушы",
        "card_number": "Сигналдық карточканың нөмірі",
        "formed_at": "Қалыптастырылған күні",
        "card_created_at": "Карточка жасалған күн",
        "snapshot_at": "Деректердің сақталған көрінісі",
        "object_code": "Нысан коды",
        "object_name": "Нысанның атауы",
        "developer": "Құрылыс салушы",
        "bin": "БСН",
        "address": "Мекенжайы",
        "territory": "Аумақ",
        "zone": "Сигналдық аймақ (анықтамалық)",
        "score": "Ұпай (анықтамалық)",
        "confirmed": "Сақталған көріністе расталған",
        "to_check": "Тексеруді қажет етеді",
        "no_items": "Сақталған көріністің бұл бөлімінде мәліметтер көрсетілмеген.",
        "draft": "ЖОБА · ЖІБЕРІЛМЕГЕН",
        "demo_stamp": "ДЕМО-ДЕРЕКТЕР",
        "neutral_note": (
            "Мәліметтер карточканың сақталған көрінісінен алынған және тексеруге арналған. "
            "Бұл құжат құқық бұзушылықты немесе кінәні анықтамайды әрі карточканың немесе "
            "қосымшалардың жіберілгенін өздігінен растамайды. Сигналдық аймақ пен ұпай тек анықтамалық көрсеткіштер."
        ),
        "outgoing_number": "Шығыс нөмірі",
        "signatory": "Қолы / аты-жөні",
        "date": "Күні",
        "appendices": "Қосымшалар тізімі",
        "appendix_no": "№",
        "appendix_name": "Атауы",
        "appendix_pages": "Парақтар",
        "no_data": "Сақталған көріністе дерек жоқ",
        "publication_sheet": "Жарияланымдар",
        "publication_meta_sheet": "Мәліметтер",
        "publication_title": "Сигналдық карточканың сақталған көрінісіндегі жарияланымдар",
        "ad_id": "Жарияланым ID-і",
        "platform": "Дереккөз / платформа",
        "account": "Аккаунт",
        "published_at": "Жарияланған күні",
        "captured_at": "Түсірілген күні",
        "url": "Сақталған көріністегі URL",
        "ad_status": "Сақталған көріністегі жарияланым мәртебесі",
        "fragments": "Тіркелген сөз тіркестері",
        "review_status": "Сақталған көріністегі тексеру мәртебесі",
        "appendix_refs": "Қосымшаға сілтемелер",
        "demo": "Демо-деректер",
        "account_missing": "Сақталған көріністе жазылмаған",
        "human_confirmed": "Адам растаған",
        "human_rejected": "Адам қабылдамаған",
        "ai_hypothesis": "ЖИ гипотезасы — адам растамаған",
        "unreviewed": "Сақталған көріністе адамның жеке тексеруі белгіленбеген",
        "fact_kind_note": "«Факт» жіктемесі адамның растағанын өздігінен білдірмейді.",
        "snapshot_only": "Тек карточканың сақталған көрінісіндегі жарияланымдар берілген.",
        "metadata": "Параметр",
        "value": "Мәні",
        "row_count": "Жарияланым жолдарының саны",
        "card_state": "Карточканың күйі",
        "draft_state": "ЖОБА · ЖІБЕРІЛМЕГЕН",
        "not_draft_state": "Берілген деректерде жоба ретінде белгіленбеген",
        "package_note": "Бұл файл карточканың немесе материалдар пакетінің жіберілгенін өздігінен растамайды.",
        "snapshot_source": "Жолдардың дереккөзі",
        "account_source": "Дереккөз аккаунты",
        "source_note": "Дереккөз аккаунты сақталған көріністе жоқ.",
        "review_rule": "Тексеру мәртебесінің ережесі",
        "dispatch_status": "Жіберу туралы мәлімет",
        "empty_rows": "Сақталған көріністе жарияланым жоқ.",
        "inventory_title": "Қосымшалар тізімдемесі",
        "inventory_intro": "Сигналдық карточкаға қосымшалар.",
        "page": "Бет",
        "inventory_file": "Файл атауы",
        "inventory_label": "Сипаттамасы",
        "inventory_source": "Дереккөз",
        "inventory_id": "Дәлел ID-і",
        "inventory_type": "Файл түрі",
        "inventory_size": "Өлшемі, байт",
        "inventory_sha": "SHA-256",
        "inventory_capture": "Түсірілген уақыты",
        "inventory_url": "Дереккөз / URL",
        "inventory_included": "Пакетке қосылуы",
        "included": "Қосылған",
        "omitted": "Қосылмаған",
        "omission_title": "Қосылмаған және қолжетімсіз файлдар",
        "omission_reason": "Себебі / ескертпе",
        "no_omissions": "Тіркелген ерекшелік жоқ.",
        "missing_metadata": "Көрсетілмеген",
        "self_exclusion": "Бұл тізімдемеде оның өз өлшемі немесе хеші көрсетілмейді.",
        "inventory_note": (
            "Тізімдеме берілген метадеректерді көрсетеді және файлдардың құқықтық маңызын, "
            "құқық бұзушылықты немесе кінәні анықтамайды."
        ),
    },
}

_POST_STATUS = {
    "ru": {"captured": "Собрана", "verification": "На проверке", "removed_saved": "Удалена, копия сохранена"},
    "kk": {"captured": "Жиналған", "verification": "Тексерілуде", "removed_saved": "Жойылған, көшірмесі сақталған"},
}


def _labels(locale: str) -> dict[str, str]:
    return _TEXT.get(locale, _TEXT["ru"])


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    return _INVALID_XML.sub("", str(value))


def _display_datetime(value: object, labels: Mapping[str, str]) -> str:
    missing = labels["missing_metadata"]
    if value is None or value == "":
        return missing
    parsed: datetime | None = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    else:
        try:
            parsed = datetime.fromisoformat(_clean_text(value).replace("Z", "+00:00"))
        except ValueError:
            return _clean_text(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    parsed = parsed.astimezone(UTC)
    return f"{parsed:%d.%m.%Y %H:%M} UTC"


def _utc_datetime(value: datetime | str) -> datetime:
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def _items(value: object) -> Sequence[Mapping[str, Any]]:
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return tuple(item for item in value if isinstance(item, Mapping))
    return ()


def render_cover_docx(data: Mapping[str, Any]) -> bytes:
    """Create an editable, neutral cover from the card snapshot descriptor."""
    locale = _clean_text(data["locale"])
    labels = _labels(locale)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    card_number = _clean_text(data["card_number"])
    if bool(data["is_draft"]):
        watermark = document.add_paragraph()
        watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = watermark.add_run(labels["draft"])
        run.bold = True
        run.font.size = Pt(12)
    if bool(data["is_demo"]):
        demo = document.add_paragraph()
        demo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = demo.add_run(labels["demo_stamp"])
        run.bold = True

    title = document.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(f"{labels['cover']}\n{labels['title']} № {card_number}")
    recipient = document.add_paragraph()
    recipient.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    recipient.add_run(f"{labels['recipient']}: {_clean_text(data['recipient_label'])}").bold = True

    details = (
        (labels["formed_at"], _display_datetime(data["formed_at"], labels)),
        (labels["card_created_at"], _display_datetime(data["card_created_at"], labels)),
        (labels["snapshot_at"], _display_datetime(data["snapshot_at"], labels)),
        (labels["object_code"], _clean_text(data["object_code"]) or labels["no_data"]),
        (labels["object_name"], _clean_text(data["object_name"]) or labels["no_data"]),
        (labels["developer"], _clean_text(data.get("developer_name")) or labels["no_data"]),
        (labels["bin"], _clean_text(data.get("bin")) or labels["no_data"]),
        (labels["address"], _clean_text(data.get("address")) or labels["no_data"]),
        (labels["territory"], _clean_text(data["territory_label"]) or labels["no_data"]),
    )
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in details:
        cells = table.add_row().cells
        cells[0].text = _clean_text(key)
        cells[1].text = _clean_text(value)
    zone = _clean_text(data.get("zone_label"))
    score = data.get("score")
    if zone or score is not None:
        context = zone or labels["no_data"]
        if score is not None:
            context = f"{context}; {labels['score'].lower()}: {_clean_text(score)}"
        document.add_paragraph(f"{labels['zone']}: {context}.")

    document.add_paragraph(_clean_text(labels["neutral_note"]))
    _add_fact_list(document, labels["confirmed"], data.get("verified_facts"), labels)
    _add_fact_list(document, labels["to_check"], data.get("to_check_facts"), labels)

    document.add_heading(labels["outgoing_number"], level=2)
    document.add_paragraph(f"{labels['outgoing_number']}: ______________________________")
    document.add_paragraph(f"{labels['signatory']}: ______________________________________")
    document.add_paragraph(f"{labels['date']}: ____________________________________________")
    document.add_heading(labels["appendices"], level=2)
    appendix_table = document.add_table(rows=1, cols=3)
    appendix_table.style = "Table Grid"
    headings = (labels["appendix_no"], labels["appendix_name"], labels["appendix_pages"])
    for cell, text in zip(appendix_table.rows[0].cells, headings, strict=True):
        cell.text = text
    appendix_names = data.get("appendices")
    appendix_names = appendix_names if isinstance(appendix_names, list) else []
    for index, appendix_name in enumerate(appendix_names, start=1):
        filename = _clean_text(appendix_name)
        if not filename:
            continue
        cells = appendix_table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = filename
        cells[2].text = "________"

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_fact_list(document: Any, heading: str, facts: object, labels: Mapping[str, str]) -> None:
    document.add_heading(heading, level=2)
    strings = [_clean_text(item) for item in facts if item is not None] if isinstance(facts, list) else []
    if not strings:
        document.add_paragraph(labels["no_items"])
        return
    for item in strings:
        document.add_paragraph(item, style="List Bullet")


def _hit_review(hit: Mapping[str, Any], labels: Mapping[str, str]) -> str:
    confirmed = hit.get("confirmed")
    if confirmed is True:
        return labels["human_confirmed"]
    if confirmed is False:
        return labels["human_rejected"]
    if hit.get("kind") == "ai_hypothesis":
        return labels["ai_hypothesis"]
    return labels["unreviewed"]


def render_publications_xlsx(
    locale: str,
    rows: Sequence[Mapping[str, Any]],
    *,
    card_number: str,
    is_draft: bool,
    generated_at: datetime | str,
) -> bytes:
    """Build a literal-string XLSX from only the publications in the supplied card snapshot."""
    labels = _labels(locale)
    export_rows: list[tuple[str, ...]] = []
    for publication in rows:
        hits = _items(publication.get("hits"))
        fragments = []
        review = []
        for hit in hits:
            fragment = _clean_text(hit.get("fragment"))
            if fragment:
                fragments.append(f"«{fragment}»")
            review.append(f"{fragment or labels['no_data']}: {_hit_review(hit, labels)}")
        evidence_ids = publication.get("evidence_ids")
        references = [f"E-{_clean_text(value)}" for value in evidence_ids] if isinstance(evidence_ids, list) else []
        status = _clean_text(publication.get("status"))
        localized_status = _POST_STATUS.get(locale, _POST_STATUS["ru"]).get(status, status or labels["no_data"])
        export_rows.append(
            (
                _clean_text(publication.get("ad_post_id")) or labels["no_data"],
                _clean_text(publication.get("platform")) or labels["no_data"],
                labels["account_missing"],
                _display_datetime(publication.get("published_at"), labels),
                _display_datetime(publication.get("captured_at"), labels),
                _clean_text(publication.get("url")) or labels["no_data"],
                localized_status,
                "\n".join(fragments) or labels["no_data"],
                "\n".join(review) or labels["no_data"],
                "; ".join(references) or labels["no_data"],
                labels["demo"] if publication.get("is_demo") else "",
            )
        )

    workbook_metadata = (
        (labels["card_number"], _clean_text(card_number)),
        (labels["card_state"], labels["draft_state"] if is_draft else labels["not_draft_state"]),
        (labels["snapshot_source"], labels["snapshot_only"]),
        (labels["account_source"], labels["source_note"]),
        (labels["review_rule"], labels["fact_kind_note"]),
        (labels["dispatch_status"], labels["package_note"]),
        (labels["row_count"], str(len(export_rows))),
    )
    columns = (
        ExportColumn(labels["ad_id"], 16),
        ExportColumn(labels["platform"], 20),
        ExportColumn(labels["account"], 26),
        ExportColumn(labels["published_at"], 20),
        ExportColumn(labels["captured_at"], 20),
        ExportColumn(labels["url"], 42),
        ExportColumn(labels["ad_status"], 24),
        ExportColumn(labels["fragments"], 48),
        ExportColumn(labels["review_status"], 52),
        ExportColumn(labels["appendix_refs"], 26),
        ExportColumn(labels["demo"], 16),
    )
    return build_workbook(
        sheet_name=labels["publication_sheet"],
        metadata_sheet_name=labels["publication_meta_sheet"],
        columns=columns,
        rows=export_rows,
        metadata=workbook_metadata,
        export_title=labels["publication_title"],
        generated_at=_utc_datetime(generated_at),
        generated_at_label=labels["formed_at"],
        metadata_headers=(labels["metadata"], labels["value"]),
    )


def _inventory_context(
    locale: str,
    rows: Sequence[Mapping[str, Any]],
    *,
    card_number: str,
    is_draft: bool,
    generated_at: datetime | str,
) -> dict[str, Any]:
    labels = _labels(locale)
    attachments = []
    omissions = []
    for row in rows:
        included = bool(row.get("included"))
        item = {
            "filename": _clean_text(row.get("filename")) or labels["missing_metadata"],
            "label": _clean_text(row.get("label")) or labels["missing_metadata"],
            "source_label": _clean_text(row.get("source_label")) or labels["missing_metadata"],
            "evidence_id": _clean_text(row.get("evidence_id")) or labels["missing_metadata"],
            "content_type": _clean_text(row.get("content_type")) or labels["missing_metadata"],
            "size_bytes": _clean_text(row.get("size_bytes")) or labels["missing_metadata"],
            "sha256": _clean_text(row.get("sha256")) or labels["missing_metadata"],
            "captured_at": _display_datetime(row.get("captured_at"), labels),
            "source_url": _clean_text(row.get("source_url")) or labels["missing_metadata"],
            "included": labels["included"] if included else labels["omitted"],
            "note": _clean_text(row.get("note")) or ("" if included else labels["missing_metadata"]),
        }
        attachments.append(item)
        if not included:
            omissions.append(item)
    return {
        "t": labels,
        "locale": locale if locale in _TEXT else "ru",
        "card_number": _clean_text(card_number),
        "generated_at": _display_datetime(generated_at, labels),
        "is_draft": is_draft,
        "attachments": attachments,
        "omissions": omissions,
    }


def _render_inventory_document(context: Mapping[str, Any]) -> PDFDocument:
    html = _env.get_template("signal_material_inventory.html").render(**context)
    return HTML(string=html, base_url=str(TEMPLATES_DIR)).render()


def render_inventory_pdf(
    locale: str,
    rows: Sequence[Mapping[str, Any]],
    *,
    card_number: str,
    is_draft: bool,
    generated_at: datetime | str,
) -> bytes:
    """Render a metadata-only appendix inventory; no evidence bytes are opened or embedded."""
    context = _inventory_context(locale, rows, card_number=card_number, is_draft=is_draft, generated_at=generated_at)
    return bytes(_render_inventory_document(context).write_pdf())
