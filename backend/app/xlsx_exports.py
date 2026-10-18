"""Small, safe helpers for staff XLSX exports."""

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from re import compile
from typing import Any
from urllib.parse import quote, urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_EXPORT_ROWS = 10_000
_INVALID_XML = compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass(frozen=True)
class ExportColumn:
    title: str
    width: float
    number_format: str | None = None


@dataclass(frozen=True)
class LinkCell:
    label: str
    url: str


def _text(value: str) -> str:
    """Keep user content literal in the workbook and remove XML 1.0 control characters."""
    return _INVALID_XML.sub("", value)


def _write_value(cell: Any, value: Any, number_format: str | None = None) -> None:
    if isinstance(value, LinkCell):
        parts = urlsplit(value.url)
        if parts.scheme not in {"http", "https"} or not parts.netloc or parts.username or parts.password:
            raise ValueError("XLSX hyperlinks must use an absolute HTTP(S) URL")
        cell.value = _text(value.label)
        cell.data_type = "s"
        cell.hyperlink = value.url
        cell.font = Font(color="0563C1", underline="single")
    elif isinstance(value, str):
        cell.value = _text(value)
        # openpyxl treats leading '=' as a formula unless the cell type is forced back to text.
        cell.data_type = "s"
    else:
        cell.value = value
    if number_format and value is not None:
        cell.number_format = number_format
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(
    *,
    sheet_name: str,
    metadata_sheet_name: str,
    columns: Sequence[ExportColumn],
    rows: Iterable[Sequence[Any]],
    metadata: Sequence[tuple[str, Any]],
    export_title: str,
    generated_at: datetime,
    generated_at_label: str,
    metadata_headers: tuple[str, str],
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "SAQ"
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise RuntimeError("New workbook has no active worksheet")
    sheet.title = sheet_name
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    bottom_border = Border(bottom=Side(style="thin", color="8EAADB"))
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        _write_value(cell, column.title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = bottom_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.row_dimensions[1].height = 32

    last_row = 1
    for last_row, values in enumerate(rows, start=2):
        if len(values) != len(columns):
            raise ValueError("XLSX row does not match the header width")
        for index, (column, value) in enumerate(zip(columns, values, strict=True), start=1):
            _write_value(sheet.cell(row=last_row, column=index), value, column.number_format)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(last_row, 1)}"

    info = workbook.create_sheet(metadata_sheet_name)
    info.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    title_cell = info.cell(row=1, column=1)
    _write_value(title_cell, export_title)
    title_cell.font = Font(size=14, bold=True, color="17365D")
    _write_value(info.cell(row=2, column=1), generated_at_label)
    _write_value(info.cell(row=2, column=2), generated_at.replace(tzinfo=None), "dd.mm.yyyy hh:mm")
    _write_value(info.cell(row=4, column=1), metadata_headers[0])
    _write_value(info.cell(row=4, column=2), metadata_headers[1])
    for cell in info[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = bottom_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_index, (label, value) in enumerate(metadata, start=5):
        _write_value(info.cell(row=row_index, column=1), label)
        _write_value(info.cell(row=row_index, column=2), value)
    info.freeze_panes = "A5"
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 90
    for row in info.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _trusted_origin(value: str | None, allowed_origins: Sequence[str], *, origin_only: bool) -> str | None:
    if not value:
        return None
    try:
        parts = urlsplit(value)
    except ValueError:
        return None
    try:
        if (
            parts.scheme not in {"http", "https"}
            or not parts.netloc
            or not parts.hostname
            or (parts.port is not None and not 0 < parts.port < 65_536)
            or parts.username
            or parts.password
            or (origin_only and (parts.path not in {"", "/"} or parts.query or parts.fragment))
        ):
            return None
    except ValueError:
        return None
    origin = f"{parts.scheme}://{parts.netloc}"
    return origin if origin in allowed_origins else None


def app_origin(
    request_origin: str | None,
    allowed_origins: Sequence[str],
    *,
    referer: str | None = None,
    request_base_url: str | None = None,
) -> str | None:
    """Choose a same-site origin for links, accepting browser context only from the CORS allowlist."""
    candidates = (
        (request_origin, True),
        (referer, False),
        (request_base_url, False),
        *((origin, True) for origin in allowed_origins),
    )
    for candidate, origin_only in candidates:
        origin = _trusted_origin(candidate, allowed_origins, origin_only=origin_only)
        if origin:
            return origin
    return None


def content_disposition(filename: str, ascii_fallback: str) -> str:
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename, safe='')}"


def utc_datetime(value: datetime | None) -> datetime | None:
    """Return an Excel-compatible naive UTC timestamp; Excel cannot store timezone-aware datetimes."""
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).replace(tzinfo=None)
