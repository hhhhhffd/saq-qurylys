"""In-memory formats for a signal-card materials package; fixtures contain invented data only."""

import re
from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocxDocument
from openpyxl import load_workbook

from app.modules.signals import material_documents

FORMED_AT = datetime(2026, 10, 17, 12, 30, tzinfo=UTC)


def _cover(**overrides: object) -> dict[str, object]:
    values: dict[str, object] = {
        "locale": "ru",
        "card_number": "SG-2026-0042",
        "formed_at": FORMED_AT.isoformat(),
        "card_created_at": "2026-10-16T08:00:00+00:00",
        "snapshot_at": "2026-10-16T08:00:00+00:00",
        "recipient_label": "Областной департамент (тест)",
        "object_code": "OBJ-0042",
        "object_name": "ЖК «Көктем»",
        "address": "Казахстан, тестовый адрес",
        "developer_name": "ТОО «Үлгі»",
        "bin": "000000000000",
        "territory_label": "Тестовая область",
        "zone_label": "Жёлтая",
        "score": 42,
        "is_demo": False,
        "is_draft": True,
        "verified_facts": ["Проверка документа подтверждена в сохранённом снимке."],
        "to_check_facts": ["Фрагмент публикации требует проверки."],
    }
    values.update(overrides)
    return values


def _document_text(document: DocxDocument) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def _page_texts(document: object) -> list[str]:
    pages = []
    for page in document.pages:  # type: ignore[attr-defined]
        parts = [box.text for box in page._page_box.descendants() if getattr(box, "text", None)]
        pages.append(re.sub(r"\s+", " ", " ".join(parts)))
    return pages


def test_cover_docx_is_editable_neutral_and_contains_snapshot_context() -> None:
    content = material_documents.render_cover_docx(_cover())
    assert content.startswith(b"PK")
    with ZipFile(BytesIO(content)) as archive:
        assert "word/document.xml" in archive.namelist()
        assert "word/styles.xml" in archive.namelist()
    document = Document(BytesIO(content))
    text = _document_text(document)

    for expected in (
        "SG-2026-0042",
        "ТОО «Үлгі»",
        "Казахстан, тестовый адрес",
        "Подтверждено в снимке",
        "Проверка документа подтверждена",
        "Требует проверки",
        "Фрагмент публикации требует проверки",
        "ЧЕРНОВИК · НЕ НАПРАВЛЕНО",
        "Исходящий номер",
        "Подпись / Ф. И. О.",
        "Дата",
        "Перечень приложений",
        "не устанавливает факт нарушения",
        "не подтверждает отправку карточки",
    ):
        assert expected in text


def test_cover_docx_lists_only_supplied_appendices() -> None:
    data = _cover(appendices=["SAQ_card.pdf", "evidence-70.png"])
    document = Document(BytesIO(material_documents.render_cover_docx(data)))
    appendix_table = document.tables[-1]

    assert len(appendix_table.rows) == 3
    assert [row.cells[1].text for row in appendix_table.rows[1:]] == ["SAQ_card.pdf", "evidence-70.png"]
    assert all(row.cells[2].text for row in appendix_table.rows[1:])


def test_demo_cover_has_no_demo_stamp() -> None:
    document = Document(BytesIO(material_documents.render_cover_docx(_cover(is_demo=True))))

    assert "ДЕМО" not in _document_text(document)


def test_cover_docx_supports_kazakh_and_strips_invalid_xml_controls() -> None:
    data = _cover(
        locale="kk",
        card_number="KZ-0007\x00",
        recipient_label="Тест алушы",
        address="Тест мекенжайы\x0b",
        is_draft=False,
        verified_facts=[],
        to_check_facts=[],
    )
    document = Document(BytesIO(material_documents.render_cover_docx(data)))
    text = _document_text(document)

    assert "KZ-0007" in text
    assert "ЖОБА · ЖІБЕРІЛМЕГЕН" not in text
    assert "Ілеспе хат" in text
    assert "Тексеруді қажет етеді" in text
    assert "Тест мекенжайы" in text
    assert "\x00" not in text
    assert "\x0b" not in text


def test_publications_xlsx_uses_only_given_snapshot_rows_and_literal_safe_cells() -> None:
    rows = [
        {
            "ad_post_id": 17,
            "platform": "instagram",
            "url": '=HYPERLINK("https://example.invalid")\x00',
            "published_at": "2026-10-12T10:00:00+00:00",
            "captured_at": "2026-10-13T11:15:00+00:00",
            "status": "captured",
            "is_demo": False,
            "evidence_ids": [70, 71],
            "hits": [
                {"fragment": "Брондау — баға!", "kind": "fact", "confirmed": True},
                {"fragment": "Төлем жоспары", "kind": "ai_hypothesis", "confirmed": None},
                {"fragment": "Сөздік сәйкестігі", "kind": "fact", "confirmed": None},
                {"fragment": "Бұрын расталған тұжырым", "kind": "fact", "confirmed": False},
            ],
        }
    ]
    content = material_documents.render_publications_xlsx(
        "kk", rows, card_number="SG-2026-0042", is_draft=True, generated_at=FORMED_AT
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Жарияланымдар"]
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and cell.column is not None:
            headers[cell.value] = cell.column
    assert sheet.max_row == 2
    assert sheet.auto_filter.ref == "A1:J2"
    assert sheet.freeze_panes == "A2"
    assert sheet.cell(row=2, column=headers["Жарияланым ID-і"]).value == "17"
    assert sheet.cell(row=2, column=headers["Дереккөз / платформа"]).value == "instagram"
    assert sheet.cell(row=2, column=headers["Аккаунт"]).value == "Сақталған көріністе жазылмаған"
    assert sheet.cell(row=2, column=headers["Жарияланған күні"]).value == "12.10.2026 10:00 UTC"
    assert sheet.cell(row=2, column=headers["Түсірілген күні"]).value == "13.10.2026 11:15 UTC"
    assert sheet.cell(row=2, column=headers["Сақталған көріністегі URL"]).value == (
        '=HYPERLINK("https://example.invalid")'
    )
    assert sheet.cell(row=2, column=headers["Қосымшаға сілтемелер"]).value == "E-70; E-71"
    fragments = sheet.cell(row=2, column=headers["Тіркелген сөз тіркестері"]).value
    review = sheet.cell(row=2, column=headers["Сақталған көріністегі тексеру мәртебесі"]).value
    assert isinstance(fragments, str)
    assert isinstance(review, str)
    assert "Брондау — баға!" in fragments
    assert "Төлем жоспары" in fragments
    assert "Адам растаған" in review
    assert "Адам қабылдамаған" in review
    assert "ЖИ гипотезасы" in review
    assert "адамның жеке тексеруі белгіленбеген" in review
    metadata_sheet = workbook["Мәліметтер"]
    metadata = {
        metadata_sheet.cell(row=row, column=1).value: metadata_sheet.cell(row=row, column=2).value
        for row in range(5, metadata_sheet.max_row + 1)
    }
    assert metadata_sheet.cell(row=2, column=2).value == FORMED_AT.replace(tzinfo=None)
    assert metadata["Сигналдық карточканың нөмірі"] == "SG-2026-0042"
    assert metadata["Карточканың күйі"] == "ЖОБА · ЖІБЕРІЛМЕГЕН"
    review_rule = metadata["Тексеру мәртебесінің ережесі"]
    assert isinstance(review_rule, str)
    assert "«Факт» жіктемесі адамның растағанын" in review_rule
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)
    assert all(cell.data_type == "s" for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert all("\x00" not in str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)


def test_inventory_pdf_lists_metadata_omissions_and_excludes_its_own_hash() -> None:
    rows = [
        {
            "filename": "proof-70.png",
            "label": "Скриншот публикации",
            "source_label": "instagram",
            "evidence_id": 70,
            "content_type": "image/png",
            "size_bytes": 2048,
            "sha256": "a" * 64,
            "captured_at": "2026-10-13T11:15:00+00:00",
            "source_url": "https://example.invalid/post/17",
            "included": True,
            "note": None,
        },
        {
            "filename": "proof-71.mhtml",
            "label": "Архив страницы",
            "source_label": "instagram",
            "evidence_id": 71,
            "content_type": "multipart/related",
            "size_bytes": None,
            "sha256": None,
            "captured_at": None,
            "source_url": None,
            "included": False,
            "note": "Файл недоступен",
        },
    ]
    content = material_documents.render_inventory_pdf(
        "ru", rows, card_number="SG-2026-0042", is_draft=True, generated_at=FORMED_AT
    )
    assert content.startswith(b"%PDF-")
    context = material_documents._inventory_context(
        "ru", rows, card_number="SG-2026-0042", is_draft=True, generated_at=FORMED_AT
    )
    pages = _page_texts(material_documents._render_inventory_document(context))
    text = " ".join(pages)

    for expected in (
        "Опись приложений № SG-2026-0042",
        "ЧЕРНОВИК · НЕ НАПРАВЛЕНО",
        "proof-70.png",
        "Скриншот публикации",
        "proof-71.mhtml",
        "Файл недоступен",
        "Эта опись не содержит собственного размера или хеша.",
        "не устанавливает правовое значение файлов",
    ):
        assert expected in text
    assert "a" * 64 in re.sub(r"\s+", "", text)
    assert "https://example.invalid/post/17" in re.sub(r"\s+", "", text)
