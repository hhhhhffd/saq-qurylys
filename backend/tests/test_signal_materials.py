"""Scoped preview and immutable ZIP package for signal-card materials."""

from __future__ import annotations

import hashlib
from datetime import UTC, date, datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.document import Document as WordDocument
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import storage
from app.models import AuditLog, Evidence
from app.models.enums import EvidenceKind, Recipient, SignalStatus
from app.modules.signals import materials, pdf
from app.storage import EvidenceLink
from tests.conftest import World, auth, make_object
from tests.support_risk import NOW, add_ad, png_bytes
from tests.test_signals_pdf import build_card


def _word_text(document: WordDocument) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def test_materials_preview_is_staff_only_and_uses_card_scope(
    client: TestClient, db: Session, world: World, bucket: str
) -> None:
    view, evidence_ids = build_card(db, world, ads=2)
    url = f"/api/v1/signals/{view.card.id}/materials"

    assert client.get(url).status_code == 401
    assert client.get(url, headers=auth(world.users["gask"])).status_code == 403
    assert client.get(url, headers=auth(world.users["mio_semey"])).status_code == 404

    response = client.get(url, headers=auth(world.users["mio_ust"]))

    assert response.status_code == 200
    payload = response.json()
    assert {item["kind"] for item in payload["files"] if item["kind"] != "evidence"} == {
        "cover_docx",
        "signal_pdf",
        "publications_xlsx",
        "inventory_pdf",
    }
    listed_evidence = {item["evidence_id"] for item in payload["files"] if item["kind"] == "evidence"}
    assert listed_evidence == set(evidence_ids["scan"] + evidence_ids["screenshots"])
    assert all(item["available"] and item["selected"] for item in payload["files"] if item["kind"] == "evidence")
    assert payload["is_draft"] is True
    assert payload["warnings"] == []
    assert payload["limits"] == {
        "max_evidence_files": 50,
        "max_file_bytes": 20 * 1024 * 1024,
        "max_total_bytes": 60 * 1024 * 1024,
    }


def test_materials_preview_excludes_same_territory_other_object_publication_evidence(
    client: TestClient, db: Session, world: World, bucket: str
) -> None:
    view, _ = build_card(db, world, ads=1)
    other_object = make_object(db, world.ust, "ЖК Другой")
    other_ad = add_ad(db, other_object, NOW)
    other_evidence = storage.put_evidence(
        db,
        png_bytes(),
        EvidenceKind.screenshot_png,
        "image/png",
        link=EvidenceLink(ad_post_id=other_ad.id),
    )

    sections = dict(view.card.sections)
    publications = list(sections["publications"])
    publications[0] = {
        **publications[0],
        "ad_post_id": other_ad.id,
        "evidence_ids": [other_evidence.id],
    }
    sections["publications"] = publications
    view.card.sections = sections
    db.flush()

    response = client.get(
        f"/api/v1/signals/{view.card.id}/materials",
        headers=auth(world.users["mio_ust"]),
    )

    assert response.status_code == 200
    files = response.json()["files"]
    assert other_evidence.id not in {item["evidence_id"] for item in files}
    assert any(
        str(other_ad.id) in item["source_label"] and item["reason"] == materials._labels("ru")["snapshot_missing"]
        for item in response.json()["missing_materials"]
    )


def test_zip_contains_real_documents_and_exact_evidence_without_dispatch_mutation(
    client: TestClient, db: Session, world: World, bucket: str
) -> None:
    view, evidence_ids = build_card(db, world, ads=2)
    response = client.get(f"/api/v1/signals/{view.card.id}/package.zip", headers=auth(world.users["mio_ust"]))

    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert names
        assert all(name.startswith("SAQ_") and ".." not in name.split("/") for name in names)
        contents = {name.rsplit("/", 1)[-1]: archive.read(name) for name in names}

    assert any(name.endswith(".docx") for name in contents)
    assert any(name.endswith(".xlsx") for name in contents)
    assert sum(data.startswith(b"%PDF-") for data in contents.values()) == 2
    cover_name = next(name for name in contents if name.endswith(".docx"))
    cover = Document(BytesIO(contents[cover_name]))
    cover_text = _word_text(cover)
    assert "Материалы к сигнальной карточке" in cover_text
    assert all(name in cover_text for name in set(contents) - {cover_name})
    workbook = load_workbook(
        BytesIO(next(data for name, data in contents.items() if name.endswith(".xlsx"))), read_only=True
    )
    assert workbook.sheetnames
    workbook.close()

    attached_ids = set(evidence_ids["scan"] + evidence_ids["screenshots"])
    evidence_by_id = {row.id: row for row in db.scalars(select(Evidence).where(Evidence.id.in_(attached_ids)))}
    included = {name for name in contents if name.startswith("evidence-")}
    assert len(included) == len(attached_ids)
    for name in included:
        evidence_id = int(name.removeprefix("evidence-").split(".", 1)[0])
        assert hashlib.sha256(contents[name]).hexdigest() == evidence_by_id[evidence_id].sha256

    db.refresh(view.card)
    assert view.card.status is SignalStatus.requires_confirmation
    assert view.card.sent_at is None
    audit_entry = db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "signal_card",
            AuditLog.entity_id == view.card.id,
            AuditLog.action == "signal.materials_package.download",
        )
    )
    assert audit_entry is not None


def test_package_selection_is_explicit_and_foreign_ids_share_one_error(
    client: TestClient,
    db: Session,
    world: World,
    bucket: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    view, evidence_ids = build_card(db, world, ads=1)
    user_headers = auth(world.users["mio_ust"])
    url = f"/api/v1/signals/{view.card.id}/package.zip"
    monkeypatch.setattr(materials, "render_cover_docx", lambda data: b"DOCX")
    monkeypatch.setattr(materials, "render_publications_xlsx", lambda *args, **kwargs: b"XLSX")
    monkeypatch.setattr(materials.pdf, "render_pdf", lambda *args, **kwargs: b"%PDF-draft")
    monkeypatch.setattr(materials, "render_inventory_pdf", lambda *args, **kwargs: b"%PDF-inventory")

    empty = client.get(url, headers=user_headers, params={"include_evidence": "false"})
    assert empty.status_code == 200
    with ZipFile(BytesIO(empty.content)) as archive:
        assert not any(name.rsplit("/", 1)[-1].startswith("evidence-") for name in archive.namelist())

    selected_id = evidence_ids["screenshots"][0]
    subset = client.get(url, headers=user_headers, params=[("evidence_id", str(selected_id))])
    assert subset.status_code == 200
    with ZipFile(BytesIO(subset.content)) as archive:
        evidence_names = [name for name in archive.namelist() if "/evidence-" in name]
    assert len(evidence_names) == 1
    assert evidence_names[0].endswith(f"evidence-{selected_id}.png")

    invalid = client.get(url, headers=user_headers, params={"evidence_id": "999999999"})
    assert invalid.status_code == 422
    assert invalid.json()["detail"]["code"] == "invalid_evidence_selection"
    mixed = client.get(
        url,
        headers=user_headers,
        params=[("include_evidence", "false"), ("evidence_id", str(selected_id))],
    )
    assert mixed.status_code == 422
    assert mixed.json()["detail"]["code"] == "invalid_evidence_selection"

    monkeypatch.setattr(materials, "MAX_PACKAGE_EVIDENCE_FILES", 1)
    too_many = client.get(url, headers=user_headers, params={"include_evidence": "true"})
    assert too_many.status_code == 413
    assert too_many.json()["detail"]["code"] == "package_evidence_limit"


def test_inventory_marks_unavailable_files_as_missing_not_employee_excluded(
    client: TestClient,
    db: Session,
    world: World,
    bucket: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    view, _ = build_card(db, world, ads=1)
    inventory_rows: list[dict[str, object]] = []
    cover_data: dict[str, object] = {}

    def capture_inventory(locale: str, rows: list[dict[str, object]], **kwargs: object) -> bytes:
        inventory_rows.extend(rows)
        return b"%PDF-inventory"

    def capture_cover(data: dict[str, object]) -> bytes:
        cover_data.update(data)
        return b"DOCX"

    monkeypatch.setattr(materials, "_storage_available", lambda evidence: False)
    monkeypatch.setattr(materials, "render_cover_docx", capture_cover)
    monkeypatch.setattr(materials, "render_publications_xlsx", lambda *args, **kwargs: b"XLSX")
    monkeypatch.setattr(materials.pdf, "render_pdf", lambda *args, **kwargs: b"%PDF-draft")
    monkeypatch.setattr(materials, "render_inventory_pdf", capture_inventory)

    response = client.get(f"/api/v1/signals/{view.card.id}/package.zip", headers=auth(world.users["mio_ust"]))

    assert response.status_code == 200
    unavailable_rows = [row for row in inventory_rows if row.get("evidence_id") is not None]
    assert unavailable_rows
    assert all(row["included"] is False for row in unavailable_rows)
    assert all(row["note"] == "Файл недоступен в хранилище" for row in unavailable_rows)
    with ZipFile(BytesIO(response.content)) as archive:
        contents = {name.rsplit("/", 1)[-1]: archive.read(name) for name in archive.namelist()}
    cover_name = next(name for name in contents if name.endswith(".docx"))
    appendices = cover_data["appendices"]
    assert isinstance(appendices, list)
    assert set(appendices) == set(contents) - {cover_name}
    missing_names = {str(row["filename"]) for row in unavailable_rows}
    assert not any(name in appendices for name in missing_names)


def test_sent_card_uses_archived_pdf_bytes_and_never_rerenders(
    client: TestClient, db: Session, world: World, bucket: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    view, _ = build_card(db, world, ads=0)
    immutable_pdf = b"%PDF-1.4\narchived card\n%%EOF"
    archive = storage_put_pdf(db, immutable_pdf, view.card.id)
    view.card.status = SignalStatus.sent
    view.card.recipient = Recipient.mio
    view.card.sent_at = datetime.now(UTC)
    view.card.due_date = date.today()
    db.flush()
    monkeypatch.setattr(materials.pdf, "render_pdf", lambda *args, **kwargs: pytest.fail("sent PDF rerendered"))

    response = client.get(f"/api/v1/signals/{view.card.id}/package.zip", headers=auth(world.users["mio_ust"]))

    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive_file:
        pdf_files = [name for name in archive_file.namelist() if "/02_card_" in name]
        assert len(pdf_files) == 1
        assert archive_file.read(pdf_files[0]) == immutable_pdf
    assert archive.sha256 == hashlib.sha256(immutable_pdf).hexdigest()
    db.refresh(view.card)
    assert view.card.status is SignalStatus.sent
    assert view.card.sent_at is not None


def test_sent_card_without_archive_is_reported_without_rerender(
    client: TestClient,
    db: Session,
    world: World,
    bucket: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    view, _ = build_card(db, world, ads=0)
    view.card.status = SignalStatus.sent
    view.card.recipient = Recipient.mio
    view.card.sent_at = datetime.now(UTC)
    view.card.due_date = date.today()
    db.flush()
    monkeypatch.setattr(materials.pdf, "render_pdf", lambda *args, **kwargs: pytest.fail("sent PDF rerendered"))
    inventory_rows: list[dict[str, object]] = []

    def capture_inventory(locale: str, rows: list[dict[str, object]], **kwargs: object) -> bytes:
        inventory_rows.extend(rows)
        return b"%PDF-inventory"

    monkeypatch.setattr(materials, "render_cover_docx", lambda data: b"DOCX")
    monkeypatch.setattr(materials, "render_publications_xlsx", lambda *args, **kwargs: b"XLSX")
    monkeypatch.setattr(materials, "render_inventory_pdf", capture_inventory)

    response = client.get(f"/api/v1/signals/{view.card.id}/package.zip", headers=auth(world.users["mio_ust"]))

    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        assert not any("/02_card_" in name for name in archive.namelist())
    assert any("сохранённая копия" in str(row.get("note", "")).lower() for row in inventory_rows)


def storage_put_pdf(db: Session, payload: bytes, card_id: int) -> Evidence:
    return storage.put_evidence(
        db,
        payload,
        EvidenceKind.pdf,
        "application/pdf",
        link=EvidenceLink(signal_card_id=card_id),
    )


def test_draft_pdf_has_warning_and_only_reads_allowlisted_evidence(db: Session, world: World, bucket: str) -> None:
    view, _ = build_card(db, world, ads=1)

    context = pdf.build_context(
        db,
        view,
        "ru",
        draft=True,
        allowed_evidence_ids=set(),
        max_embedded_total_bytes=0,
    )
    html = pdf.render_html(context)

    assert "ЧЕРНОВИК · НЕ НАПРАВЛЕНО" in html
    assert "data:image" not in html


def test_pdf_image_embedding_stops_after_actual_byte_limit(
    db: Session, world: World, bucket: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, ids = build_card(db, world, ads=1)
    evidence = db.get(Evidence, ids["screenshots"][0])
    assert evidence is not None
    evidence.size_bytes = 1

    class FakeResponse:
        def __init__(self) -> None:
            self.read_sizes: list[int] = []
            self.closed = False
            self.released = False

        def read(self, size: int) -> bytes:
            self.read_sizes.append(size)
            return b"12345"[:size]

        def close(self) -> None:
            self.closed = True

        def release_conn(self) -> None:
            self.released = True

    fake = FakeResponse()
    monkeypatch.setattr(pdf.storage, "open_evidence", lambda row: fake)
    image, note, size = pdf._embedded_image(evidence, pdf.get_labels("ru"), max_bytes=4)

    assert image is None
    assert note == pdf.get_labels("ru")["not_embedded"]
    assert size == 0
    assert fake.read_sizes == [5]
    assert fake.closed and fake.released


def test_evidence_reader_uses_actual_bounded_sizes_for_package_limits(
    db: Session, world: World, bucket: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, ids = build_card(db, world, ads=1)
    evidence = db.get(Evidence, ids["screenshots"][0])
    assert evidence is not None

    class FakeResponse:
        def __init__(self, data: bytes) -> None:
            self.data = data
            self.read_sizes: list[int] = []
            self.closed = False
            self.released = False

        def read(self, size: int) -> bytes:
            self.read_sizes.append(size)
            return self.data[:size]

        def close(self) -> None:
            self.closed = True

        def release_conn(self) -> None:
            self.released = True

    file_overflow = FakeResponse(b"12345")
    monkeypatch.setattr(materials, "MAX_PACKAGE_FILE_BYTES", 4)
    monkeypatch.setattr(materials.storage, "open_evidence", lambda row: file_overflow)
    with pytest.raises(HTTPException) as file_error:
        materials._read_limited(evidence, 100, "ru")
    assert file_error.value.status_code == 413
    assert isinstance(file_error.value.detail, dict)
    assert file_error.value.detail.get("code") == "package_file_too_large"
    assert file_overflow.read_sizes == [5]
    assert file_overflow.closed and file_overflow.released

    total_overflow = FakeResponse(b"12345")
    monkeypatch.setattr(materials, "MAX_PACKAGE_FILE_BYTES", 10)
    monkeypatch.setattr(materials.storage, "open_evidence", lambda row: total_overflow)
    with pytest.raises(HTTPException) as total_error:
        materials._read_limited(evidence, 4, "ru")
    assert total_error.value.status_code == 413
    assert isinstance(total_error.value.detail, dict)
    assert total_error.value.detail.get("code") == "package_total_size_limit"
    assert total_overflow.read_sizes == [5]
    assert total_overflow.closed and total_overflow.released
