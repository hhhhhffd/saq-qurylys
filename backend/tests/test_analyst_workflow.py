"""Synthetic analyst workflow from collector intake to a filtered signal package."""

from __future__ import annotations

import hashlib
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import seed, storage
from app.models import AdPost, ConstructionObject, Evidence, PhraseHit, SignalCard, SourceAccount
from app.models.enums import (
    AdOrigin,
    AdStatus,
    EvidenceKind,
    PhraseCategory,
    Platform,
    SignalStatus,
    Zone,
)
from app.modules.nlp import service as nlp
from app.modules.risk import service as risk
from app.storage import EvidenceLink
from app.userbot import service as userbot
from tests.conftest import World, auth
from tests.support_risk import png_bytes

ADS = "/api/v1/ads"
SIGNALS = "/api/v1/signals"
DEMO_CASES = {
    "ЖК «Северный бриз»": "yellow",
    "ЖК «Золотая осень»": "red",
    "ЖК «Тихая гавань»": "green",
}


def _seed_control_cases(db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, ConstructionObject]:
    """Use the existing fictional A/B/C seed records without their extra demo publications."""
    source_dir = seed.SEED_DATA_DIR
    payload: dict[str, Any] = json.loads((source_dir / "demo_objects.json").read_text(encoding="utf-8"))
    selected = [item for item in payload["objects"] if item["name"] in DEMO_CASES]
    assert {item["name"] for item in selected} == set(DEMO_CASES)
    for item in selected:
        item["ads"] = []
    developer_bins = {item["developer_bin"] for item in selected}
    payload["objects"] = selected
    payload["developers"] = [developer for developer in payload["developers"] if developer["bin"] in developer_bins]

    seed.seed_reference(db)
    seed.load_districts(db)
    seed.seed_demo_users(db)
    fixture_dir = tmp_path / "demo-seed"
    fixture_dir.mkdir()
    (fixture_dir / "demo_objects.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(seed, "SEED_DATA_DIR", fixture_dir)
    ids = seed.seed_demo_objects(db)
    objects = {obj.name: obj for obj in db.scalars(select(ConstructionObject).where(ConstructionObject.id.in_(ids)))}
    assert set(objects) == set(DEMO_CASES)
    assert all(obj.is_demo for obj in objects.values())
    return objects


def test_demo_analyst_workflow_reaches_filtered_signal_package(
    client: TestClient,
    db: Session,
    world: World,
    bucket: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    objects = _seed_control_cases(db, tmp_path, monkeypatch)
    queued: list[tuple[int, bool, int | None]] = []

    def record_pipeline(ad_id: int, *, capture: bool, requested_by: int | None) -> bool:
        queued.append((ad_id, capture, requested_by))
        return True

    monkeypatch.setattr(userbot, "dispatch_pipeline", record_pipeline)
    account = SourceAccount(
        platform=Platform.telegram,
        handle="analyst_workflow_demo",
        territory_id=world.ust.id,
    )
    db.add(account)
    db.flush()

    posts = [
        ("ЖК «Северный бриз»: платная бронь квартир", objects["ЖК «Северный бриз»"]),
        ("ЖК «Золотая осень»: внести депозит 100 000 ₸", objects["ЖК «Золотая осень»"]),
        ("ЖК «Золотая осень»: бронь квартиры без оплаты", objects["ЖК «Золотая осень»"]),
        ("ЖК «Тихая гавань»: обычный обзор готового двора", objects["ЖК «Тихая гавань»"]),
    ]
    now = datetime.now(UTC)
    phrases = userbot.load_phrases(db)
    ads: list[AdPost] = []
    evidence_bytes: dict[int, bytes] = {}
    for index, (text, obj) in enumerate(posts, start=1):
        post = userbot.build_post(
            username=account.handle,
            message_id=index,
            text=text,
            links=[],
            published_at=now - timedelta(hours=index),
            has_photo=False,
            is_channel=True,
            channel_title="Synthetic demo channel",
            phrases=phrases,
        )
        assert post is not None
        ad_id = userbot.store_message(db, account.id, post, [], capture=False)
        assert ad_id is not None
        ad = db.get(AdPost, ad_id)
        assert ad is not None and ad.origin is AdOrigin.collector
        ad.is_demo = True
        nlp.analyze_ad(db, ad)
        nlp.link_ad(db, ad)
        assert ad.status is AdStatus.verification
        assert ad.object_id == obj.id
        image = png_bytes(color=("#2255aa", "#aa5522", "#228855", "#7755aa")[index - 1])
        evidence = storage.put_evidence(
            db,
            image,
            EvidenceKind.screenshot_png,
            "image/png",
            source_url=post.url,
            link=EvidenceLink(ad_post_id=ad.id),
        )
        evidence_bytes[evidence.id] = image
        ads.append(ad)
    db.flush()

    assert len(queued) == 4 and {item[0] for item in queued} == {ad.id for ad in ads}
    initial_queue = client.get(ADS, headers=auth(world.users["mio_ust"]), params={"queue": True, "is_demo": True})
    assert initial_queue.status_code == 200
    queue_rows = initial_queue.json()["items"]
    assert {row["id"] for row in queue_rows} == {ad.id for ad in ads}
    assert all(row["is_demo"] and row["origin"] == "collector" for row in queue_rows)

    headers = auth(world.users["mio_ust"])
    for ad, (_, obj) in zip(ads, posts, strict=True):
        linked = client.patch(f"{ADS}/{ad.id}", headers=headers, json={"object_id": obj.id})
        assert linked.status_code == 200, linked.text
        assert linked.json()["link_method"] == "manual" and linked.json()["status"] == "captured"

    def hit_for(ad: AdPost, category: PhraseCategory) -> PhraseHit:
        return db.scalars(select(PhraseHit).where(PhraseHit.ad_post_id == ad.id, PhraseHit.category == category)).one()

    decisions = ((ads[1], PhraseCategory.deposit, True), (ads[2], PhraseCategory.booking, False))
    for ad, category, confirmed in decisions:
        hit = hit_for(ad, category)
        response = client.post(
            f"{ADS}/{ad.id}/hits/{hit.id}/confirm",
            headers=headers,
            json={"confirmed": confirmed},
        )
        assert response.status_code == 200, response.text
        assert response.json()["confirmed"] is confirmed

    deferred = client.patch(
        f"{ADS}/{ads[0].id}/triage",
        headers=headers,
        json={"state": "deferred", "note": "Вернуться после проверки источника"},
    )
    assert deferred.status_code == 200 and deferred.json()["review_state"] == "deferred"
    deferred_rows = client.get(
        ADS, headers=headers, params={"queue": True, "is_demo": True, "review_state": "deferred"}
    )
    assert {row["id"] for row in deferred_rows.json()["items"]} == {ads[0].id}
    assert ads[0].id not in {
        row["id"] for row in client.get(ADS, headers=headers, params={"queue": True, "is_demo": True}).json()["items"]
    }

    reopened = client.patch(f"{ADS}/{ads[0].id}/triage", headers=headers, json={"state": "unreviewed"})
    assert reopened.status_code == 200 and reopened.json()["review_state"] == "unreviewed"
    next_rows = client.get(ADS, headers=headers, params={"queue": True, "is_demo": True})
    assert ads[0].id in {row["id"] for row in next_rows.json()["items"]}

    for ad in (ads[1], ads[2], ads[3]):
        reviewed = client.patch(f"{ADS}/{ad.id}/triage", headers=headers, json={"state": "reviewed"})
        assert reviewed.status_code == 200, reviewed.text
    reviewed_rows = client.get(
        ADS, headers=headers, params={"queue": True, "is_demo": True, "review_state": "reviewed"}
    )
    assert {row["id"] for row in reviewed_rows.json()["items"]} == {ads[1].id, ads[2].id, ads[3].id}

    results = {name: risk.recalculate(db, obj.id) for name, obj in objects.items()}
    zones = {name: result.risk_score.zone for name, result in results.items() if result is not None}
    assert zones == {name: Zone(zone) for name, zone in DEMO_CASES.items()}
    red_obj = objects["ЖК «Золотая осень»"]
    red_result = results[red_obj.name]
    assert red_result is not None and red_result.signal_card_id is not None
    card = db.get(SignalCard, red_result.signal_card_id)
    assert card is not None and card.status is SignalStatus.requires_confirmation
    assert bool(card.sections)

    artifact_dir = tmp_path / "artifacts"
    artifact_dir.mkdir()
    package_response = client.get(f"{SIGNALS}/{card.id}/package.zip", headers=headers)
    assert package_response.status_code == 200, package_response.text
    package_path = artifact_dir / "analyst-workflow-package.zip"
    package_path.write_bytes(package_response.content)
    with ZipFile(BytesIO(package_response.content)) as archive:
        names = archive.namelist()
        assert names and all(".." not in name.split("/") for name in names)
        contents = {name.rsplit("/", 1)[-1]: archive.read(name) for name in names}
    docx_bytes = next(data for name, data in contents.items() if name.endswith(".docx"))
    xlsx_bytes = next(data for name, data in contents.items() if name.endswith(".xlsx"))
    pdf_files = [data for name, data in contents.items() if name.endswith(".pdf")]
    assert len(pdf_files) == 2 and all(data.startswith(b"%PDF-") for data in pdf_files)
    assert all(data.rstrip().endswith(b"%%EOF") for data in pdf_files)
    cover = Document(BytesIO(docx_bytes))
    assert any("Материалы к сигнальной карточке" in paragraph.text for paragraph in cover.paragraphs)
    workbook = load_workbook(BytesIO(xlsx_bytes), read_only=True)
    assert workbook.sheetnames
    workbook.close()

    red_ad_ids = {ads[1].id, ads[2].id}
    red_evidence = list(db.scalars(select(Evidence).where(Evidence.ad_post_id.in_(red_ad_ids))))
    included_names = {name for name in contents if name.startswith("evidence-")}
    assert len(included_names) == len(red_evidence) == 2
    for name in included_names:
        evidence_id = int(name.removeprefix("evidence-").split(".", 1)[0])
        evidence = next(row for row in red_evidence if row.id == evidence_id)
        assert hashlib.sha256(contents[name]).hexdigest() == evidence.sha256
        assert evidence_bytes[evidence_id] == contents[name]

    export_response = client.get(
        f"{SIGNALS}/export.xlsx",
        headers=headers,
        params={"object_id": red_obj.id, "status": SignalStatus.requires_confirmation.value},
    )
    assert export_response.status_code == 200, export_response.text
    export_path = artifact_dir / "analyst-workflow-filtered-signals.xlsx"
    export_path.write_bytes(export_response.content)
    export_book = load_workbook(BytesIO(export_response.content), read_only=True)
    data_sheet = export_book[export_book.sheetnames[0]]
    assert data_sheet.max_row == 2
    assert [cell.value for cell in data_sheet[2]][2] == card.number
    export_book.close()

    before = {
        ad.id: (ad.text, ad.object_id, ad.status, ad.review_state, ad.max_phrase_weight, ad.total_phrase_weight)
        for ad in ads
    }
    db.refresh(card)
    assert card.status is SignalStatus.requires_confirmation and card.sent_at is None
    for ad in ads:
        db.refresh(ad)
        assert (
            ad.text,
            ad.object_id,
            ad.status,
            ad.review_state,
            ad.max_phrase_weight,
            ad.total_phrase_weight,
        ) == before[ad.id]
