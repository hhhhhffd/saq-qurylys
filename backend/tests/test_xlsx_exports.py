from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AuditLog, ConstructionObject, Developer, Document, SignalCard
from app.models.enums import CheckStatus, DocType, Locale, Recipient, SignalOrigin, SignalStatus, Zone
from app.xlsx_exports import app_origin
from tests.conftest import World, auth, make_object

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_export_links_use_allowlisted_referer_when_origin_is_absent() -> None:
    allowed = ["http://localhost:5173", "http://localhost:8080"]
    assert app_origin(None, allowed, referer="http://localhost:8080/signals?status=sent") == "http://localhost:8080"
    assert app_origin(None, allowed, referer="http://untrusted.example/signals") == "http://localhost:5173"


def _workbook(response: Any) -> Workbook:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    return load_workbook(BytesIO(response.content), data_only=False)


def _rows_by_header(sheet: Any) -> tuple[dict[str, int], list[dict[str, Any]]]:
    headers = {cell.value: cell.column for cell in sheet[1]}
    rows = [
        {header: sheet.cell(row=row_number, column=column).value for header, column in headers.items()}
        for row_number in range(2, sheet.max_row + 1)
    ]
    return headers, rows


def _add_signal(
    db: Session,
    obj: ConstructionObject,
    *,
    number: str,
    status: SignalStatus,
    recipient: Recipient,
    due_date: date | None,
    sent_at: datetime | None = None,
    response_at: datetime | None = None,
) -> SignalCard:
    card = SignalCard(
        number=number,
        object_id=obj.id,
        status=status,
        origin=SignalOrigin.manual,
        sections={},
        recipient=recipient,
        review_days=5,
        sent_at=sent_at,
        due_date=due_date,
        response_at=response_at,
    )
    db.add(card)
    db.flush()
    return card


def test_objects_export_uses_full_filtered_registry_and_territory_scope(
    client: TestClient, db: Session, world: World
) -> None:
    matching = [
        make_object(
            db,
            world.ust,
            f"ЖК Экспорт {index:02}",
            zone=Zone.red,
            score=100 - index,
            district_id=world.district.id,
        )
        for index in range(55)
    ]
    make_object(db, world.ust, "ЖК Экспорт зелёный", zone=Zone.green, district_id=world.district.id)
    outside = make_object(db, world.semey, "ЖК Экспорт чужой город", zone=Zone.red, district_id=None)

    response = client.get(
        "/api/v1/objects/export.xlsx",
        headers=auth(world.users["mio_ust"]),
        params={
            "zone": "red",
            "district_id": world.district.id,
            "territory_id": world.ust.id,
            "q": "ЖК Экспорт",
        },
    )

    workbook = _workbook(response)
    sheet = workbook["Объекты"]
    headers, rows = _rows_by_header(sheet)
    assert len(rows) == 55  # more than the list's default page of 50 rows
    assert {row["ID"] for row in rows} == {obj.id for obj in matching}
    assert outside.id not in {row["ID"] for row in rows}
    assert [row["Балл риска"] for row in rows] == sorted(
        (obj.score for obj in matching if obj.score is not None), reverse=True
    )
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:T56"
    object_link = sheet.cell(row=2, column=headers["Карточка"]).hyperlink
    assert object_link is not None
    assert object_link.target == f"http://localhost:8080/objects/{matching[0].id}"
    assert "filename*=UTF-8''SAQ_%D0%96%D0%9A_" in response.headers["content-disposition"]

    metadata = workbook["Об экспорте"]
    metadata_rows = {
        metadata.cell(row=r, column=1).value: metadata.cell(row=r, column=2).value
        for r in range(5, metadata.max_row + 1)
    }
    assert metadata_rows["Строк выгружено"] == 55
    assert metadata_rows["Зона риска"] == "Красная"
    assert metadata_rows["Территория"] == world.ust.name

    audit_row = db.scalars(select(AuditLog).where(AuditLog.action == "objects.export")).one()
    assert audit_row.after["row_count"] == 55
    assert audit_row.after["territory_id"] == world.ust.id


def test_objects_export_is_valid_when_empty_and_rejects_forbidden_roles(client: TestClient, world: World) -> None:
    response = client.get("/api/v1/objects/export.xlsx?q=missing", headers=auth(world.users["mio_ust"]))
    workbook = _workbook(response)
    assert workbook["Объекты"].max_row == 1
    metadata = workbook["Об экспорте"]
    assert {
        metadata.cell(row=r, column=1).value: metadata.cell(row=r, column=2).value
        for r in range(5, metadata.max_row + 1)
    }["Строк выгружено"] == 0
    assert {
        metadata.cell(row=r, column=1).value: metadata.cell(row=r, column=2).value
        for r in range(5, metadata.max_row + 1)
    }["Результат"] == "Нет данных по выбранным фильтрам"
    assert client.get("/api/v1/objects/export.xlsx", headers=auth(world.users["gask"])).status_code == 403


def test_objects_export_localizes_and_keeps_formula_like_user_values_as_text(
    client: TestClient, db: Session, world: World
) -> None:
    world.users["mio_ust"].locale = Locale.kk
    obj = make_object(
        db,
        world.ust,
        '=HYPERLINK("https://example.invalid","x")',
        address="\t=1+1",
        zone=Zone.yellow,
        is_demo=True,
    )

    response = client.get("/api/v1/objects/export.xlsx", headers=auth(world.users["mio_ust"]))

    workbook = _workbook(response)
    sheet = workbook["Объектілер"]
    headers, rows = _rows_by_header(sheet)
    assert "Атауы" in headers
    assert workbook.sheetnames[1] == "Экспорт туралы"
    name_cell = sheet.cell(row=2, column=headers["Атауы"])
    address_cell = sheet.cell(row=2, column=headers["Мекенжайы"])
    assert rows[0]["ID"] == obj.id
    assert isinstance(name_cell.value, str)
    assert name_cell.value.startswith("=HYPERLINK(") and name_cell.data_type == "s"
    assert isinstance(address_cell.value, str)
    assert address_cell.value == "\t=1+1" and address_cell.data_type == "s"
    assert rows[0]["Демо-деректер"] == "Иә"


def test_objects_export_reports_limit_without_truncating(
    client: TestClient, db: Session, world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_object(db, world.ust, "ЖК Лимит один")
    make_object(db, world.ust, "ЖК Лимит два")
    monkeypatch.setattr("app.modules.objects.router.MAX_EXPORT_ROWS", 1)

    response = client.get("/api/v1/objects/export.xlsx", headers=auth(world.users["mio_ust"]))

    assert response.status_code == 413
    assert response.json()["detail"]["code"] == "too_many_export_rows"


def test_object_export_uses_shared_legal_basis_status_and_real_dates(
    client: TestClient, db: Session, world: World
) -> None:
    developer = Developer(bin="123456789012", name="ТОО Экспорт")
    db.add(developer)
    db.flush()
    checked_at = datetime(2026, 9, 21, 12, 30, tzinfo=UTC)
    obj = make_object(
        db,
        world.ust,
        "ЖК Основание",
        developer_id=developer.id,
        zone=Zone.green,
        score=10,
        last_scored_at=checked_at,
        updated_at=checked_at,
    )
    db.add_all(
        [
            Document(
                object_id=obj.id,
                doc_type=DocType.guarantee,
                check_status=CheckStatus.confirmed,
                checked_at=checked_at,
            ),
            Document(
                object_id=obj.id, doc_type=DocType.dsu_permit, check_status=CheckStatus.not_found, checked_at=checked_at
            ),
        ]
    )
    db.flush()

    workbook = _workbook(client.get("/api/v1/objects/export.xlsx", headers=auth(world.users["mio_ust"])))
    headers, rows = _rows_by_header(workbook["Объекты"])

    assert rows[0]["БИН"] == developer.bin
    assert rows[0]["Правовое основание"] == "Подтверждено"
    assert rows[0]["Гарантия КЖК — проверка"] == "Подтверждено"
    assert rows[0]["Разрешение ДСУ — проверка"] == "Не найдено"
    assert rows[0]["Последняя проверка (UTC)"] == datetime(2026, 9, 21, 12, 30)
    assert rows[0]["Последний расчёт (UTC)"] == datetime(2026, 9, 21, 12, 30)
    assert workbook["Объекты"].cell(row=2, column=headers["Последний расчёт (UTC)"]).number_format == "dd.mm.yyyy hh:mm"


def test_signal_export_keeps_recipient_scope_status_filters_and_recorded_dates(
    client: TestClient, db: Session, world: World
) -> None:
    obj = make_object(db, world.ust, "ЖК Сигнал", zone=Zone.red)
    past_due = date.today() - timedelta(days=2)
    now = datetime.now(UTC)
    response_at = now - timedelta(days=1)
    sent = _add_signal(
        db,
        obj,
        number="SQ-2026-9001",
        status=SignalStatus.sent,
        recipient=Recipient.gask,
        due_date=past_due,
        sent_at=now - timedelta(days=4),
    )
    responded = _add_signal(
        db,
        obj,
        number="SQ-2026-9002",
        status=SignalStatus.responded,
        recipient=Recipient.gask,
        due_date=past_due,
        sent_at=now - timedelta(days=4),
        response_at=response_at,
    )
    _add_signal(
        db,
        obj,
        number="SQ-2026-9003",
        status=SignalStatus.sent,
        recipient=Recipient.law_enforcement,
        due_date=past_due,
        sent_at=now - timedelta(days=4),
    )
    _add_signal(
        db,
        obj,
        number="SQ-2026-9004",
        status=SignalStatus.requires_confirmation,
        recipient=Recipient.gask,
        due_date=None,
    )
    outside = make_object(db, world.semey, "ЖК Чужой город")
    _add_signal(
        db,
        outside,
        number="SQ-2026-9005",
        status=SignalStatus.sent,
        recipient=Recipient.gask,
        due_date=past_due,
        sent_at=now - timedelta(days=4),
    )

    response = client.get(
        "/api/v1/signals/export.xlsx",
        headers=auth(world.users["gask"]),
        params={"object_id": obj.id},
    )

    workbook = _workbook(response)
    sheet = workbook["Сигналы"]
    _, rows = _rows_by_header(sheet)
    assert {row["ID"] for row in rows} == {sent.id, responded.id}
    sent_row = next(row for row in rows if row["ID"] == sent.id)
    response_row = next(row for row in rows if row["ID"] == responded.id)
    assert sent_row["Просрочена на дату формирования"] == "Да"
    assert response_row["Просрочена на дату формирования"] == "Нет"
    response_date = response_row["Дата внесения ответа (UTC)"]
    assert abs(response_date - response_at.replace(tzinfo=None)) < timedelta(milliseconds=1)
    assert sent_row["Статус карточки"] == "Направлена"
    assert "Внутреннее направление (UTC)" in rows[0]
    metadata = workbook["Об экспорте"]
    meta_rows = {
        metadata.cell(row=r, column=1).value: metadata.cell(row=r, column=2).value
        for r in range(5, metadata.max_row + 1)
    }
    risk_note = meta_rows["Примечание о риске"]
    assert isinstance(risk_note, str)
    assert "не доказывает нарушение" in risk_note
    headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
    assert all(isinstance(header, str) for header in headers)
    assert "достав" not in " ".join(header.lower() for header in headers if isinstance(header, str))
    sent_row_number = next(row_number for row_number, row in enumerate(rows, start=2) if row["ID"] == sent.id)
    signal_link = sheet.cell(row=sent_row_number, column=2).hyperlink
    assert signal_link is not None
    assert signal_link.target == f"http://localhost:8080/signals/{sent.id}"

    filtered = _workbook(
        client.get(
            "/api/v1/signals/export.xlsx",
            headers=auth(world.users["gask"]),
            params={"object_id": obj.id, "status": "responded"},
        )
    )
    _, filtered_rows = _rows_by_header(filtered["Сигналы"])
    assert [row["ID"] for row in filtered_rows] == [responded.id]


def test_signal_export_is_empty_or_rejects_above_limit(
    client: TestClient, db: Session, world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    empty = _workbook(client.get("/api/v1/signals/export.xlsx", headers=auth(world.users["admin"])))
    assert empty["Сигналы"].max_row == 1
    metadata = empty["Об экспорте"]
    meta_rows = {
        metadata.cell(row=r, column=1).value: metadata.cell(row=r, column=2).value
        for r in range(5, metadata.max_row + 1)
    }
    assert meta_rows["Результат"] == "Нет данных по выбранным фильтрам"

    obj = make_object(db, world.ust)
    for index in range(2):
        _add_signal(
            db,
            obj,
            number=f"SQ-2026-910{index}",
            status=SignalStatus.requires_confirmation,
            recipient=Recipient.gask,
            due_date=None,
        )
    monkeypatch.setattr("app.modules.signals.router.MAX_EXPORT_ROWS", 1)
    limited = client.get("/api/v1/signals/export.xlsx", headers=auth(world.users["admin"]))
    assert limited.status_code == 413
    assert limited.json()["detail"]["code"] == "too_many_export_rows"
